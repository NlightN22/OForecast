from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="4472C4")
TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


class ExcelWriter:
    def save(self, results: list[dict], filename: str | Path) -> None:
        path = Path(filename)
        path.parent.mkdir(parents=True, exist_ok=True)
        totals, employees = self._tables(results)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            totals.to_excel(
                writer,
                sheet_name="Brand Totals",
                index=False,
                startrow=1,
            )
            employees.to_excel(
                writer,
                sheet_name="Employee Forecasts",
                index=False,
                startrow=1,
            )

        self._format(path)

    @staticmethod
    def _tables(results: list[dict]) -> tuple[pd.DataFrame, pd.DataFrame]:
        total_rows = []
        employee_rows = []

        for item in results:
            response = item["response"]
            intervals = response.get("intervals", [])
            interval = intervals[0] if intervals else {}
            row = {
                "Brand": item["brand"],
                "Next Period": response.get("next_period"),
                "Forecast": interval.get("point"),
                "Lower 80%": interval.get("lo80"),
                "Upper 80%": interval.get("hi80"),
                "Lower 95%": interval.get("lo95"),
                "Upper 95%": interval.get("hi95"),
                "Model": response.get("chosen_model"),
                "Dataset": response.get("chosen_dataset"),
                "Rows In": response.get("rows_in"),
                "Status": response.get("status", "ok"),
                "Error": response.get("error", ""),
            }

            if item["forecast_level"] == "brand_total":
                total_rows.append(row)
            else:
                employee_rows.append({"Employee": item["manager"], **row})

        return pd.DataFrame(total_rows), pd.DataFrame(employee_rows)

    def _format(self, path: Path) -> None:
        workbook = load_workbook(path)
        for sheet in workbook.worksheets:
            self._format_sheet(sheet)
        workbook.save(path)

    @staticmethod
    def _format_sheet(sheet) -> None:
        last_column = sheet.max_column
        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=last_column,
        )
        title = sheet.cell(row=1, column=1)
        title.value = "API version OForecast"
        title.fill = TITLE_FILL
        title.font = TITLE_FONT
        title.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 24

        for cell in sheet[2]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in sheet.iter_rows():
            for cell in row:
                cell.border = THIN_BORDER

        numeric_headers = {
            "Forecast",
            "Lower 80%",
            "Upper 80%",
            "Lower 95%",
            "Upper 95%",
            "Rows In",
        }
        for cell in sheet[2]:
            if cell.value in numeric_headers:
                for column in sheet.iter_cols(
                    min_col=cell.column,
                    max_col=cell.column,
                    min_row=3,
                ):
                    for value in column:
                        value.number_format = "#,##0.00"

        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[letter].width = min(width, 45)

        sheet.auto_filter.ref = f"A2:{get_column_letter(last_column)}{sheet.max_row}"
        sheet.freeze_panes = "A3"
