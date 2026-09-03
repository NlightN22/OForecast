"""
Save forecast results to Excel.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Border
from openpyxl.styles import Side
from openpyxl.styles import Alignment

from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="4472C4",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

TITLE_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)

THIN_BORDER = Border(

    left=Side(style="thin"),

    right=Side(style="thin"),

    top=Side(style="thin"),

    bottom=Side(style="thin"),

)


class ExcelWriter:

    def __init__(self):
        pass

    def _format(
        self,
        filename: Path,
    ):

        wb = load_workbook(filename)

        for ws in wb.worksheets:
            self._format_sheet(ws)

        wb.save(filename)

    # -------------------------------------------------------------

    def _format_sheet(self, ws) -> None:

        last_column = ws.max_column
        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=last_column,
        )
        title = ws.cell(row=1, column=1)
        title.value = "Local version OForecast"
        title.fill = TITLE_FILL
        title.font = TITLE_FONT
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        for cell in ws[2]:

            cell.fill = HEADER_FILL

            cell.font = HEADER_FONT

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            cell.border = THIN_BORDER

        for row in ws.iter_rows():

            for cell in row:

                cell.border = THIN_BORDER

        self._format_numbers(ws)

        self._format_dates(ws)

        self._autosize(ws)

        ws.auto_filter.ref = f"A2:{get_column_letter(last_column)}{ws.max_row}"
        ws.freeze_panes = "A3"

    # -------------------------------------------------------------

    @staticmethod
    def _validate(result: pd.DataFrame) -> None:
        if result.empty:
            raise ValueError("Forecast result is empty.")

    # -------------------------------------------------------------

    @staticmethod
    def _prepare_directory(path: Path) -> None:
        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

    # -------------------------------------------------------------

    @staticmethod
    def _write_forecast(
        writer: pd.ExcelWriter,
        result: pd.DataFrame,
    ) -> None:
        if "forecast_level" not in result.columns:
            return

        totals = result[result["forecast_level"] == "brand_total"]
        details = result[result["forecast_level"] == "manager_brand"]

        if not totals.empty:
            totals.to_excel(
                writer,
                sheet_name="Brand Totals",
                index=False,
                startrow=1,
            )

        if not details.empty:
            details.to_excel(
                writer,
                sheet_name="Employee Forecasts",
                index=False,
                startrow=1,
            )

        if totals.empty and details.empty:
            raise ValueError(
                "Forecast result has no rows with forecast_level "
                "'brand_total' or 'manager_brand' to write."
            )

    # -------------------------------------------------------------

    @staticmethod
    def _format_numbers(ws) -> None:
        numeric_columns = {
            "forecast",
            "lo80",
            "hi80",
            "lo95",
            "hi95",
            "rows_in",
            "rows_after_fill",
            "missing_periods_filled",
        }

        headers = {
            cell.column: str(cell.value)
            for cell in ws[2]
        }

        for column_index, header in headers.items():
            if header not in numeric_columns:
                continue

            for cell in ws.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=3,
            ):
                for item in cell:
                    item.number_format = "#,##0.00"

    # -------------------------------------------------------------

    @staticmethod
    def _format_dates(ws) -> None:
        headers = {
            cell.column: str(cell.value)
            for cell in ws[2]
        }

        for column_index, header in headers.items():
            if header != "next_period":
                continue

            for cell in ws.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=3,
            ):
                for item in cell:
                    item.alignment = Alignment(
                        horizontal="center",
                    )

    # -------------------------------------------------------------

    @staticmethod
    def _autosize(ws) -> None:
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(
                    max_length,
                    len(value),
                )

            ws.column_dimensions[column_letter].width = min(
                max_length + 2,
                60,
            )

    # -------------------------------------------------------------

    def save(
        self,
        result: pd.DataFrame,
        filename: str,
    ) -> None:
        """
        Save forecast results.
        """

        self._validate(result)

        path = Path(filename)

        self._prepare_directory(path)

        try:

            with pd.ExcelWriter(
                path,
                engine="openpyxl",
            ) as writer:

                self._write_forecast(
                    writer,
                    result,
                )

            self._format(path)

        except PermissionError:

            raise PermissionError(
                f"File '{filename}' is open in Excel. "
                "Close it and try again."
            )

        except Exception as exc:

            raise RuntimeError(
                f"Excel save error:\n{exc}"
            ) from exc
